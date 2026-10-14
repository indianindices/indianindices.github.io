import json
from calendar import monthrange
from datetime import date
from math import isfinite
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
XLSX_PATH = ROOT / "data" / "sectoraldata.xlsx"
OUT_PATH = ROOT.parent / "frontend" / "data.json"
QUARTER_ENDS = [(3, 31), (6, 30), (9, 30), (12, 31)]
METADATA_SHEET = "_RefreshStatus"
MONTHLY_SHEET = "_MonthlyReturns"
MONTHLY_METADATA_SHEET = "_MonthlyRefreshStatus"
METHODOLOGY = {"source": "Trendlyne", "returnType": "unverified", "roundingDecimals": 2, "reconciliationTolerancePp": 0.1}


def pct(value):
    return round(value * 100, 2) if isinstance(value, (int, float)) and not isinstance(value, bool) and isfinite(value) else None


def main():
    today = date.today()
    assets = {}
    refresh = {}
    monthly_assets = {}
    monthly_refresh = {}
    workbook = load_workbook(XLSX_PATH, read_only=True)
    try:
        for sheet in workbook.worksheets:
            if sheet.title in {METADATA_SHEET, MONTHLY_METADATA_SHEET}:
                records = {name: {"lastSuccess": last_success, "lastAttempt": last_attempt, "status": status}
                           for name, last_success, last_attempt, status in sheet.iter_rows(min_row=2, max_col=4, values_only=True) if name}
                if sheet.title == METADATA_SHEET:
                    refresh = records
                else:
                    monthly_refresh = records
                continue
            if sheet.title == MONTHLY_SHEET:
                for name, label, *values in sheet.iter_rows(min_row=2, max_col=15, values_only=True):
                    if not name or not (label == "5Yr Avg" or isinstance(label, int)):
                        continue
                    months = [pct(value) for value in values[:12]]
                    if isinstance(label, int):
                        months = [value if date(label, month, monthrange(label, month)[1]) < today else None for month, value in enumerate(months, 1)]
                    monthly_assets.setdefault(name, []).append({"year": None if label == "5Yr Avg" else label, "avg": label == "5Yr Avg", "m": months, "annual": pct(values[12])})
                continue
            rows = []
            for label, *values in sheet.iter_rows(min_row=2, max_col=6, values_only=True):
                if label == "5Yr Avg":
                    rows.append({"year": None, "avg": True, "q": [pct(v) for v in values[:4]], "annual": pct(values[4])})
                elif isinstance(label, int):
                    quarters = [pct(v) for v in values[:4]]
                    if label >= today.year:
                        # Source reports unfinished quarters as 0; hide them instead.
                        quarters = [v if date(label, m, d) < today else None for v, (m, d) in zip(quarters, QUARTER_ENDS)]
                    rows.append({"year": label, "avg": False, "q": quarters, "annual": pct(values[4])})
            if not any(not row["avg"] for row in rows):
                # Refuse to publish if a fetch silently returned nothing.
                raise SystemExit(f"No data for '{sheet.title}' in {XLSX_PATH.name}; aborting export.")
            assets[sheet.title] = rows
    finally:
        workbook.close()
    OUT_PATH.write_text(json.dumps({"generated": today.isoformat(), "currentYear": today.year, "assets": assets, "refresh": refresh, "monthlyAssets": monthly_assets, "monthlyRefresh": monthly_refresh, "methodology": METHODOLOGY}, allow_nan=False))
    print(f"Exported {len(assets)} assets to {OUT_PATH}")


if __name__ == "__main__":
    main()
