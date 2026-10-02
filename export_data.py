import json
from datetime import date
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).parent
XLSX_PATH = ROOT / "sectoraldata.xlsx"
OUT_PATH = ROOT / "data.json"
QUARTER_ENDS = [(3, 31), (6, 30), (9, 30), (12, 31)]


def pct(value):
    return round(value * 100, 2) if isinstance(value, (int, float)) else None


def main():
    today = date.today()
    assets = {}
    for sheet in load_workbook(XLSX_PATH, read_only=True).worksheets:
        rows = []
        for label, *values in sheet.iter_rows(min_row=2, max_col=6, values_only=True):
            if label == "5Yr Avg":
                rows.append({"year": None, "avg": True, "q": [pct(v) for v in values[:4]], "annual": pct(values[4])})
            elif isinstance(label, int):
                quarters = [pct(v) for v in values[:4]]
                if label == today.year:
                    # Source reports unfinished quarters as 0; hide them instead.
                    quarters = [v if date(label, m, d) < today else None for v, (m, d) in zip(quarters, QUARTER_ENDS)]
                rows.append({"year": label, "avg": False, "q": quarters, "annual": pct(values[4])})
        if not any(not row["avg"] for row in rows):
            # Refuse to publish if a fetch silently returned nothing.
            raise SystemExit(f"No data for '{sheet.title}' in {XLSX_PATH.name}; aborting export.")
        assets[sheet.title] = rows
    OUT_PATH.write_text(json.dumps({"generated": today.isoformat(), "currentYear": today.year, "assets": assets}))
    print(f"Exported {len(assets)} assets to {OUT_PATH}")


if __name__ == "__main__":
    main()
