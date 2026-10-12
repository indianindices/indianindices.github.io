import json
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest
from unittest.mock import patch

from openpyxl import Workbook, load_workbook

import build_sectoraldata as build
import export_data as export


class RefreshTests(unittest.TestCase):
    def setUp(self):
        self.directory = TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        root = Path(self.directory.name)
        self.workbook = root / "returns.xlsx"
        self.output = root / "data.json"
        self.rows = [[2025, 10, -5, 3, 2, 9.78]]
        self.monthly_rows = [[2025, *range(1, 13), 9.78]]
        self.periods = {"quarterly": self.rows, "monthly": self.monthly_rows}
        self.build_paths = patch.multiple(build, XLSX_PATH=self.workbook, CSV_PATH=root / "returns.csv", REQUESTED=["Nifty50", "Next50"])
        self.export_paths = patch.multiple(export, XLSX_PATH=self.workbook, OUT_PATH=self.output)
        self.build_paths.start()
        self.export_paths.start()
        self.addCleanup(self.build_paths.stop)
        self.addCleanup(self.export_paths.stop)

    def test_export_labels_unverified_return_basis_and_rejects_nonfinite_values(self):
        for value in [float("nan"), float("inf"), float("-inf"), True, None]:
            self.assertIsNone(export.pct(value))
        self.assertEqual(export.pct(0), 0)
        self.assertEqual(export.pct(-0.1), -10)
        self.seed()
        export.main()
        methodology = json.loads(self.output.read_text())["methodology"]
        self.assertEqual(methodology["returnType"], "unverified")
        self.assertEqual(methodology["reconciliationTolerancePp"], 0.1)

    def test_refresh_creates_data_directory_and_exports_from_it(self):
        directory = Path(self.directory.name) / "data"
        workbook_path = directory / "sectoraldata.xlsx"
        csv_path = directory / "sectoraldata.csv"
        with patch.multiple(build, XLSX_PATH=workbook_path, CSV_PATH=csv_path), patch.object(build, "trendlyne_returns", return_value=self.periods):
            build.main()
        self.assertTrue(workbook_path.exists())
        self.assertTrue(csv_path.exists())
        with patch.object(export, "XLSX_PATH", workbook_path):
            export.main()
        data = json.loads(self.output.read_text())
        self.assertEqual(set(data["assets"]), {"Nifty50", "Next50"})
        self.assertEqual(data["monthlyAssets"]["Nifty50"][0]["m"], list(range(1, 13)))

    def seed(self, metadata=True):
        workbook = Workbook()
        workbook.remove(workbook.active)
        for name in build.REQUESTED:
            build.write_sheet(workbook, name, self.rows)
        if metadata:
            sheet = workbook.create_sheet(build.METADATA_SHEET)
            sheet.append(["Index", "Last successful refresh", "Last refresh attempt", "Status"])
            for name in build.REQUESTED:
                sheet.append([name, "2025-01-01", "2025-01-01", "fresh"])
            sheet.sheet_state = "hidden"
        workbook.save(self.workbook)
        workbook.close()

    def test_success_and_fallback_persist_without_becoming_assets(self):
        self.seed()
        with patch.object(build, "trendlyne_returns", side_effect=[self.periods, RuntimeError("source unavailable")]):
            build.main()
        export.main()
        data = json.loads(self.output.read_text())
        self.assertEqual(set(data["assets"]), {"Nifty50", "Next50"})
        self.assertEqual(data["refresh"]["Nifty50"]["status"], "fresh")
        self.assertEqual(data["refresh"]["Nifty50"]["lastSuccess"], date.today().isoformat())
        self.assertEqual(data["refresh"]["Next50"]["status"], "cached")
        self.assertEqual(data["refresh"]["Next50"]["lastSuccess"], "2025-01-01")
        self.assertEqual(data["refresh"]["Next50"]["lastAttempt"], date.today().isoformat())
        self.assertEqual(data["assets"]["Next50"][0]["annual"], 9.78)
        self.assertNotIn(build.METADATA_SHEET, build.previous_rows())
        workbook = load_workbook(self.workbook)
        build.add_cagr(workbook)
        self.assertEqual(workbook[build.METADATA_SHEET].sheet_state, "hidden")
        self.assertEqual(workbook[build.METADATA_SHEET].max_column, 4)
        workbook.close()

    def test_legacy_workbook_does_not_invent_success_dates(self):
        self.seed(metadata=False)
        export.main()
        self.assertEqual(json.loads(self.output.read_text())["refresh"], {})
        with patch.object(build, "trendlyne_returns", side_effect=RuntimeError("source unavailable")):
            build.main()
        export.main()
        records = json.loads(self.output.read_text())["refresh"]
        self.assertTrue(all(record["lastSuccess"] is None for record in records.values()))
        self.assertTrue(all(record["status"] == "cached" for record in records.values()))

    def test_average_only_fetch_is_not_a_success(self):
        self.seed()
        with patch.object(build, "trendlyne_returns", return_value={"quarterly": [["5Yr Avg", 1, 1, 1, 1, 4]], "monthly": []}):
            build.main()
        export.main()
        self.assertTrue(all(record["status"] == "cached" for record in json.loads(self.output.read_text())["refresh"].values()))

    def test_failure_without_saved_data_still_aborts_export(self):
        with patch.object(build, "trendlyne_returns", side_effect=RuntimeError("source unavailable")):
            build.main()
        with self.assertRaises(SystemExit):
            export.main()
        self.assertFalse(self.output.exists())

    def test_monthly_fetch_cache_and_independent_status(self):
        self.seed()
        with patch.object(build, "trendlyne_returns", return_value=self.periods):
            build.main()
        export.main()
        data = json.loads(self.output.read_text())
        self.assertEqual(data["monthlyAssets"]["Nifty50"][0]["m"], list(range(1, 13)))
        self.assertEqual(data["monthlyRefresh"]["Nifty50"]["status"], "fresh")
        with patch.object(build, "trendlyne_returns", return_value={"quarterly": self.rows, "monthly": []}):
            build.main()
        export.main()
        data = json.loads(self.output.read_text())
        self.assertEqual(data["refresh"]["Nifty50"]["status"], "fresh")
        self.assertEqual(data["monthlyRefresh"]["Nifty50"]["status"], "cached")
        self.assertEqual(data["monthlyRefresh"]["Nifty50"]["lastSuccess"], date.today().isoformat())
        self.assertEqual(data["monthlyAssets"]["Nifty50"][0]["m"], list(range(1, 13)))
        self.assertEqual(set(build.previous_rows()), {"Nifty50", "Next50"})
        with patch.object(build, "trendlyne_returns", side_effect=RuntimeError("offline")):
            build.main()
        export.main()
        self.assertEqual(json.loads(self.output.read_text())["monthlyRefresh"]["Nifty50"]["status"], "cached")

    def test_monthly_export_hides_unfinished_months_including_leap_year(self):
        self.seed()
        workbook = load_workbook(self.workbook)
        sheet = workbook.create_sheet(build.MONTHLY_SHEET)
        sheet.append(["Index", "Year", *range(1, 13), "Annual"])
        sheet.append(["Nifty50", 2028, *[0.01] * 12, 0.12])
        sheet.append(["Nifty50", 2029, *[0.01] * 12, 0.12])
        workbook.save(self.workbook)
        workbook.close()
        for month, day, completed in [(2, 29, 1), (3, 1, 2), (1, 3, 0)]:
            class FixedDate(date):
                @classmethod
                def today(cls):
                    return cls(2028, month, day)

            with self.subTest(month=month, day=day), patch.object(export, "date", FixedDate):
                export.main()
                rows = json.loads(self.output.read_text())["monthlyAssets"]["Nifty50"]
                self.assertEqual(rows[0]["m"], [1.0] * completed + [None] * (12 - completed))
                self.assertEqual(rows[1]["m"], [None] * 12)

    def test_source_parses_both_periods_without_inventing_missing_values(self):
        payload = {"body": {"returnsPatternData": {
            "Quarterly": {"2025": [{"m": 3, "v": 0}, {"y": "2025", "v": 4}]},
            "Monthly": {"2025": [{"m": 1, "v": -2}, {"m": 2, "v": 0}]},
        }}}
        with patch.object(build.requests, "get") as request:
            request.return_value.json.return_value = payload
            periods = build.trendlyne_returns(1887)
        self.assertEqual(request.call_count, 1)
        self.assertEqual(periods["quarterly"], [[2025, 0, None, None, None, 4]])
        self.assertEqual(periods["monthly"], [[2025, -2, 0, *[None] * 11]])

    def test_2027_export_rolls_year_forward_and_preserves_2026_quarters(self):
        self.rows = [[2027, 4, 7, 2, 3, 4], [2026, 2, 3, 4, 5, 14.7]]
        self.seed()
        for month, expected in [(1, [None, None, None, None]), (4, [4.0, None, None, None])]:
            class FixedDate(date):
                @classmethod
                def today(cls):
                    return cls(2027, month, 3)

            with self.subTest(month=month), patch.object(export, "date", FixedDate):
                export.main()
            data = json.loads(self.output.read_text())
            self.assertEqual(data["currentYear"], 2027)
            current, previous = data["assets"]["Nifty50"]
            self.assertEqual(current["year"], 2027)
            self.assertEqual(current["q"], expected)
            self.assertEqual(previous["year"], 2026)
            self.assertEqual(previous["q"], [2.0, 3.0, 4.0, 5.0])

    def test_year_rollover_does_not_invent_a_missing_source_row(self):
        self.rows = [[2026, 2, 3, 4, 5, 14.7]]
        self.seed()

        class FixedDate(date):
            @classmethod
            def today(cls):
                return cls(2027, 1, 3)

        with patch.object(export, "date", FixedDate):
            export.main()
        data = json.loads(self.output.read_text())
        self.assertEqual(data["currentYear"], 2027)
        self.assertEqual([row["year"] for row in data["assets"]["Nifty50"]], [2026])


if __name__ == "__main__":
    unittest.main()